# -*- coding: utf-8 -*-
"""独立复算：审查《第一问详细建模说明》中的试算数字。"""
import math
import numpy as np
import tifffile
import openpyxl

BASE = 'D:/CyberStorm/Documents/数模/D题/'
DEM_PATH = BASE + '数据/镇龙乡地理空间数据/镇龙乡及周边地理数据/数字高程模型数据（DEM）/镇龙乡及周边30米DEM.tif'
NODE_PATH = BASE + '数据/无人机应急物资运输基础数据/调度中心与服务区.xlsx'
BOX_PATH = BASE + '数据/无人机应急物资运输基础数据/物资需求与配送时限.xlsx'
DRONE_PATH = BASE + '数据/无人机应急物资运输基础数据/运输无人机数据.xlsx'

# ---------- DEM ----------
dem = np.asarray(tifffile.imread(DEM_PATH), dtype=np.float64)
TL_LON, TL_LAT = 109.03277777777778, 23.224722222222223
STEP = 0.00027777777777777393
NR, NC = dem.shape

def dem_at(lon, lat):
    """最近邻取样元高程"""
    j = int(round((lon - TL_LON) / STEP))
    i = int(round((TL_LAT - lat) / STEP))
    i = min(max(i, 0), NR - 1)
    j = min(max(j, 0), NC - 1)
    return dem[i, j]

# ---------- 节点 ----------
wb = openpyxl.load_workbook(NODE_PATH, data_only=True)
ws = wb['数据']
nodes = {}
mode = None
for row in ws.iter_rows(values_only=True):
    if row and row[0] == 'O01':
        nodes['O01'] = (row[2], row[3], row[4]); mode = 'S'
    elif mode == 'S' and row and isinstance(row[0], str) and row[0].startswith('S'):
        nodes[row[0]] = (row[2], row[3], row[4])

# ---------- 机型 ----------
wb = openpyxl.load_workbook(DRONE_PATH, data_only=True)
ws = wb['数据']
hdr = None
types = {}
for row in ws.iter_rows(values_only=True):
    if row and row[0] == '机型编号':
        hdr = row
    elif hdr and row and row[0] in ('A', 'B', 'C') and isinstance(row[1], str) and len(row[1]) > 3:
        types[row[0]] = dict(zip(['id','name','m','Q','V','vc','L0','LF','E','rho','t_prep','t_load','t_hand','t_hand_box','v_up','v_dn','eta_up','eta_dn'], row))

# ---------- 货箱 ----------
wb = openpyxl.load_workbook(BOX_PATH, data_only=True)
ws = wb['逐箱货箱清单']
boxes = []
for row in ws.iter_rows(values_only=True):
    if row and isinstance(row[0], str) and '-' in row[0]:
        boxes.append(dict(id=row[0], zone=row[1], typ=row[2], w=row[3], v=row[4], first=row[5]=='是'))

# ---------- 距离 / 巡航海拔 ----------
def local_xy(lon, lat, lon0, lat0):
    """局部等距圆柱投影，以 (lon0, lat0) 为原点"""
    R = 6371000.0
    x = math.radians(lon - lon0) * R * math.cos(math.radians(lat0))
    y = math.radians(lat - lat0) * R
    return x, y

O = nodes['O01']
def segment_profile(n1, n2, n_samples=400):
    """沿线采样最高地形高程"""
    l1, a1 = nodes[n1][0], nodes[n1][1]
    l2, a2 = nodes[n2][0], nodes[n2][1]
    hs = []
    for k in range(n_samples + 1):
        t = k / n_samples
        lon = l1 + t * (l2 - l1)
        lat = a1 + t * (a2 - a1)
        hs.append(dem_at(lon, lat))
    x1, y1 = local_xy(l1, a1, O[0], O[1])
    x2, y2 = local_xy(l2, a2, O[0], O[1])
    d = math.hypot(x2 - x1, y2 - y1)
    return d, max(hs)

def L_eq(g, q):
    t = types[g]
    return t['L0'] - (t['L0'] - t['LF']) * (q / t['Q']) ** 1.5

def E_hor(g, d, q):
    return types[g]['E'] * d / L_eq(g, q)

def E_up(g, h, q):
    t = types[g]
    return (t['m'] + q) * 9.81 * h / (t['eta_up'] * 3.6e6)

def rt_energy(g, i, q, d, Hc):
    """O01->Si->O01 往返能耗"""
    hO = nodes['O01'][2]            # O01 作业高度
    hS = nodes[i][2] + 30.0         # 服务区作业高度
    # 去程: 爬升 hO->Hc, 下降 Hc->hS；返程: 爬升 hS->Hc, 下降 Hc->hO
    e = 0.0
    e += E_hor(g, d, q) + E_up(g, max(Hc - hO, 0), q)
    e += E_hor(g, d, 0) + E_up(g, max(Hc - hS, 0), 0)
    return e

def flight_time(g, i, d, Hc):
    t = types[g]
    hO = nodes['O01'][2]; hS = nodes[i][2] + 30.0
    out = max(Hc - hO, 0) / t['v_up'] + d / t['vc'] + max(Hc - hS, 0) / t['v_dn']
    back = max(Hc - hS, 0) / t['v_up'] + d / t['vc'] + max(Hc - hO, 0) / t['v_dn']
    return out + back

# ---------- 逐区计算 ----------
print('=' * 100)
print('各区水平距离 / 巡航海拔 / 最大安全载荷（二分, rho=20%）')
print('=' * 100)
results = {}
for i in sorted(nodes):
    if i == 'O01':
        continue
    d, hmax = segment_profile('O01', i)
    Hc = hmax + 50.0
    results[i] = (d, Hc)
    row = [i, f'd={d/1000:.2f}km', f'Hc={Hc:.1f}m']
    for g in ('A', 'B', 'C'):
        t = types[g]
        E_ok = t['E'] * (1 - t['rho'] / 100)
        if rt_energy(g, i, 0, d, Hc) > E_ok:
            row.append(f'{g}:不可行')
            continue
        if rt_energy(g, i, t['Q'], d, Hc) <= E_ok:
            row.append(f'{g}:{t["Q"]:.0f}kg')
            continue
        lo, hi = 0.0, t['Q']
        for _ in range(60):
            mid = (lo + hi) / 2
            if rt_energy(g, i, mid, d, Hc) <= E_ok:
                lo = mid
            else:
                hi = mid
        row.append(f'{g}:{lo:.1f}kg')
    print('  '.join(row))

# ---------- 复算 18 架次方案能耗与时间 ----------
print()
print('=' * 100)
print('复算 9C+9B 组批方案的总能耗与累计作业时间')
print('=' * 100)
zone_boxes = {}
for b in boxes:
    zone_boxes.setdefault(b['zone'], []).append(b)

def plan_energy_time(g, i, boxlist):
    d, Hc = results[i]
    q = sum(b['w'] for b in boxlist)
    t = types[g]
    E = rt_energy(g, i, q, d, Hc)
    T = t['t_prep'] + t['t_load'] * len(boxlist) + flight_time(g, i, d, Hc) + t['t_hand'] + t['t_hand_box'] * len(boxlist)
    soc = 100 * (1 - E / t['E'])
    return E, T, q, soc

def pick(zone, typ, n):
    return [b for b in zone_boxes[zone] if b['typ'] == typ][:n]

total_E, total_T = 0.0, 0.0
rows = []
# S001: 2x C (4水+1食+1卫+2医 | 剩余)
z = zone_boxes['S001']
b1 = pick('S001', '饮用水', 4) + pick('S001', '应急食品', 1) + pick('S001', '生活卫生用品', 1) + pick('S001', '医疗物资', 2)
b2 = [b for b in z if b not in b1]
for bl in (b1, b2):
    E, T, q, soc = plan_energy_time('C', 'S001', bl)
    total_E += E; total_T += T
    rows.append(('S001', 'C', q, sum(b['v'] for b in bl), len(bl), E, T, soc))
# S002, S003: C(3水+1食+1卫) + B(1水+1食+1医)
for zi in ('S002', 'S003'):
    bc = pick(zi, '饮用水', 3) + pick(zi, '应急食品', 1) + pick(zi, '生活卫生用品', 1)
    bb = pick(zi, '饮用水', 1)[1:] + pick(zi, '应急食品', 1)[1:] + pick(zi, '医疗物资', 1)
    bb = [b for b in zone_boxes[zi] if b not in bc]
    for g, bl in (('C', bc), ('B', bb)):
        E, T, q, soc = plan_energy_time(g, zi, bl)
        total_E += E; total_T += T
        rows.append((zi, g, q, sum(b['v'] for b in bl), len(bl), E, T, soc))
# S004-S008: 1x C 整区
for zi in ('S004', 'S005', 'S006', 'S007', 'S008'):
    bl = zone_boxes[zi]
    E, T, q, soc = plan_energy_time('C', zi, bl)
    total_E += E; total_T += T
    rows.append((zi, 'C', q, sum(b['v'] for b in bl), len(bl), E, T, soc))
# S009-S015: 1x B 整区
for zi in ('S009', 'S010', 'S011', 'S012', 'S013', 'S014', 'S015'):
    bl = zone_boxes[zi]
    E, T, q, soc = plan_energy_time('B', zi, bl)
    total_E += E; total_T += T
    rows.append((zi, 'B', q, sum(b['v'] for b in bl), len(bl), E, T, soc))

print(f"{'区':<6}{'机型':<4}{'质量kg':<8}{'体积m3':<8}{'箱数':<5}{'能耗kWh':<9}{'时间s':<8}{'返航SOC%':<8}")
for r in rows:
    print(f"{r[0]:<6}{r[1]:<4}{r[2]:<8.0f}{r[3]:<8.3f}{r[4]:<5}{r[5]:<9.3f}{r[6]:<8.0f}{r[7]:<8.1f}")
print('-' * 60)
print(f'总计: {len(rows)} 架次, 总能耗 {total_E:.2f} kWh, 累计作业时间 {total_T:.0f} s ({total_T/1e4:.2f} 万秒)')

# ---------- S004 临界余量 ----------
d, Hc = results['S004']
q4 = sum(b['w'] for b in zone_boxes['S004'])
E4 = rt_energy('C', 'S004', q4, d, Hc)
crit = 1 - E4 / types['C']['E']
print()
print(f'S004 整区 {q4} kg 单架 C 能耗 {E4:.3f} kWh -> 临界返航余量 {crit*100:.1f}%')

# 18 架次下界的另外核验
print()
print('各区总质量:', {z: sum(b["w"] for b in bl) for z, bl in sorted(zone_boxes.items())})
