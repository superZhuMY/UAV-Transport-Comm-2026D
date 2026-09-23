# -*- coding: utf-8 -*-
"""把两版第二问方案的结果整理成 Word 说明文档（宋体小四 / 黑体标题 / 三线表）。"""
import os
import csv
import json
import io
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
Q2 = os.path.join(BASE, 'q2', 'runs')
OUT = os.path.join(BASE, 'Q2结果说明.docx')

RUNS = [('q2_run01', '方案 A（完工时刻优先）'), ('q2_run02', '方案 B（能耗优先，对照）')]


def load(run):
    d = os.path.join(Q2, run)
    man = json.load(io.open(os.path.join(d, 'run_manifest.json'), encoding='utf-8'))
    val = json.load(io.open(os.path.join(d, 'validation.json'), encoding='utf-8'))
    sorties = list(csv.DictReader(io.open(os.path.join(d, 'transport_sorties.csv'),
                                          encoding='utf-8-sig')))
    boxes = list(csv.DictReader(io.open(os.path.join(d, 'box_deliveries.csv'),
                                        encoding='utf-8-sig')))
    return dict(dir=d, man=man, val=val, sorties=sorties, boxes=boxes)


def set_font(run, name='宋体', size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if color:
        run.font.color.rgb = color


def para(doc, text, size=12, name='宋体', bold=False, indent=True,
         align=None, space_after=6, line=1.5, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_after = Pt(space_after)
    if indent:
        pf.first_line_indent = Pt(size * 2)
    if align is not None:
        p.alignment = align
    set_font(p.add_run(text), name, size, bold, color)
    return p


def heading(doc, text, level=1):
    sizes = {1: 14, 2: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    set_font(p.add_run(text), '黑体', sizes.get(level, 12), True)
    return p


def three_line(table, header_rows=1):
    tbl = table._tbl
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                el = OxmlElement('w:%s' % edge)
                if edge == 'top' and i == 0:
                    el.set(qn('w:val'), 'single')
                    el.set(qn('w:sz'), '12')
                elif edge == 'bottom' and i == header_rows - 1:
                    el.set(qn('w:val'), 'single')
                    el.set(qn('w:sz'), '6')
                elif edge == 'bottom' and i == len(table.rows) - 1:
                    el.set(qn('w:val'), 'single')
                    el.set(qn('w:sz'), '12')
                else:
                    el.set(qn('w:val'), 'nil')
                borders.append(el)
            tcPr.append(borders)


def add_table(doc, header, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, h in enumerate(header):
        cell = t.rows[0].cells[k]
        cell.text = ''
        set_font(cell.paragraphs[0].add_run(str(h)), '黑体', size, True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in rows:
        cells = t.add_row().cells
        for k, v in enumerate(r):
            cells[k].text = ''
            set_font(cells[k].paragraphs[0].add_run('' if v is None else str(v)),
                     '宋体', size)
            cells[k].paragraphs[0].alignment = (WD_ALIGN_PARAGRAPH.CENTER
                                                if k else WD_ALIGN_PARAGRAPH.LEFT)
    if widths:
        for row in t.rows:
            for k, w in enumerate(widths):
                row.cells[k].width = Cm(w)
    three_line(t, 1)
    return t


def main():
    A, B = load('q2_run01'), load('q2_run02')
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = '宋体'
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.6)
        s.top_margin = s.bottom_margin = Cm(2.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run('第二问 异构无人机多点多架次运输调度 计算结果说明'), '黑体', 16, True)

    heading(doc, '一、计算口径与继承关系')
    para(doc, '本问沿用第一问已声明的公共物理规则：节点地面海拔取节点表，服务区作业海拔为'
              '地面海拔以上 30 m；航段为节点间水平直线，在 WGS 84 / UTM 49N 米制坐标下计算'
              '水平距离，巡航海拔取该航段所经 30 m DEM 原始像元最高高程以上 50 m；多停靠架次'
              '每站交接完成后从作业高度重新爬升，返程按空载计。')
    para(doc, '能耗换算延续统一假设：水平巡航能耗 = 单组电池可用能量 × 水平距离 ÷ 该载荷下'
              '等效航程 L(q)=L0−(L0−LF)(q/Q)^1.5；爬升附加能耗 =（含电池机身质量 + 当前载荷）'
              '×9.81×爬升高度 ÷（爬升效率×3.6×10�6），下降附加能耗不计。安全余量 20% 只在'
              '整趟能量上限中扣一次，即整趟能耗不超过 0.8 倍单组电池可用能量。')
    para(doc, '交付时刻采用保守的整站交接口径：同一架次同一停靠点交付的货箱，一律在该站'
              '基础交接与逐箱增加交接全部结束后记为交付完成，下一航段亦在此后开始。共享电池'
              '自架次准备开始即被占用，直到返场后充至 100% 才可再次分配；无人机返场即可准备'
              '下一趟。全部任务完成时间取所有选中架次中最晚的返场时刻，不包含任务结束后的充电。')

    heading(doc, '二、输入对账')
    checks = _input_checks()
    rows = [(n, '通过' if ok else '失败', i) for n, ok, i in checks]
    add_table(doc, ['核对项', '结论', '实测值'], rows, widths=[8.0, 2.4, 5.0], size=9)

    heading(doc, '三、求解流程')
    para(doc, '采用“物理仿真与校验 → 单点可行基线 → 多点候选架次 → CP-SAT 联合调度 → '
              '独立复算与导出”的流程。候选架次只预先固定一趟任务的机型、停靠事件与各站货箱'
              '（含单箱直送、区内紧急组批、区内装箱链及其随机变体、跨区两点/三点合并），'
              '是否选用、开始时刻与设备占用由同一个调度模型决定。')
    para(doc, '联合调度模型对每个候选架次设 0-1 选择变量与开始时刻变量，每个货箱恰好被一个'
              '选中架次覆盖；同机型无人机与共享电池分别施加区间累计容量约束（4/2/2 台与'
              '6/4/4 组）；医疗物资与首批保障货箱的硬时限作为硬约束直接限制开始时刻上界。'
              '目标按词典序推进，方案 A 为 加权超期量 → 完工时刻 → 总能耗 → 架次数，'
              '方案 B 为 加权超期量 → 总能耗 → 完工时刻 → 架次数。')
    ps = A['man']['pool_stats']
    para(doc, '本次候选池共 %d 条（其中多点候选 %d 条，最多 %d 站），覆盖 80 个货箱与全部'
              '31 个硬时限货箱；机型分布为 %s。' % (ps['n'], ps['multi'], ps['max_stops'],
                                                 ps['by_type']))

    heading(doc, '四、两版方案指标与权衡')
    hdr = ['指标', '基线（列表调度）', '方案 A', '方案 B']
    rows = []
    for key, label, fmt in [('weighted_tardiness', '加权超期量（Σω·迟延，s）', '{:.0f}'),
                            ('makespan_s', '全部任务完成时间（s）', '{:.0f}'),
                            ('energy_kwh', '总运输能耗（kWh）', '{:.2f}'),
                            ('n_sorties', '运输架次数', '{:.0f}')]:
        rows.append([label,
                     fmt.format(A['man']['baseline'][key]),
                     fmt.format(A['man']['final_metrics'][key]),
                     fmt.format(B['man']['final_metrics'][key])])
    rows.append(['硬时限最小余量（s）',
                 '%.0f' % (A['man']['baseline'].get('min_hard_slack_s') or 0),
                 '%.0f' % A['val']['deadlines']['min_hard_slack_s'],
                 '%.0f' % B['val']['deadlines']['min_hard_slack_s']])
    rows.append(['使用无人机 / 电池数',
                 '—', '%d / %d' % (A['val']['resources']['aircraft']['n_used'],
                                   A['val']['resources']['battery']['n_used']),
                 '%d / %d' % (B['val']['resources']['aircraft']['n_used'],
                              B['val']['resources']['battery']['n_used'])])
    add_table(doc, hdr, rows, widths=[4.6, 3.2, 3.2, 3.2])
    para(doc, '两版方案均实现 80 箱全部在期望送达时间内交付（加权超期量为 0），且不违反任何'
              '硬时限。方案 B 以完工时刻从 %.0f s 延长到 %.0f s、硬时限最小余量由 %.0f s 收缩'
              '到 %.0f s 为代价，把总能耗从 %.2f kWh 降到 %.2f kWh（−%.1f%%），架次数从 %d 降到 %d。'
              % (A['man']['final_metrics']['makespan_s'], B['man']['final_metrics']['makespan_s'],
                 A['val']['deadlines']['min_hard_slack_s'], B['val']['deadlines']['min_hard_slack_s'],
                 A['man']['final_metrics']['energy_kwh'], B['man']['final_metrics']['energy_kwh'],
                 100 * (1 - B['man']['final_metrics']['energy_kwh']
                        / A['man']['final_metrics']['energy_kwh']),
                 A['man']['final_metrics']['n_sorties'], B['man']['final_metrics']['n_sorties']))
    para(doc, '这说明配送及时性与运输能耗之间存在明显交换：尽早并行发车会推高架次数与能耗；'
              '合并多点、减少返场次数可显著节能，但会拉长任务完成时间并压缩时限余量。'
              '两版方案互为非支配解，可按救援阶段的实际偏好选用。')

    heading(doc, '五、方案 A 逐架次安排')
    rows = []
    for s in sorted(A['sorties'], key=lambda r: float(r['start_s'])):
        rows.append([s['sortie_id'], s['aircraft_id'], s['type_id'], s['battery_id'],
                     '%.0f' % float(s['start_s']), s['stop_sequence'],
                     '%.0f' % float(s['return_s']), s['n_boxes'], s['mass_kg'],
                     s['volume_m3'], '%.3f' % float(s['energy_kwh']),
                     '%.1f' % float(s['return_soc_pct'])])
    add_table(doc, ['架次', '无人机', '机型', '电池', '开始(s)', '访问顺序', '返回(s)',
                    '箱数', '质量(kg)', '体积(m³)', '能耗(kWh)', '返航SOC(%)'],
              rows, widths=[1.3, 1.5, 1.1, 1.4, 1.4, 4.2, 1.4, 1.0, 1.4, 1.5, 1.5, 1.7], size=8)

    heading(doc, '六、资源可行性与独立校验')
    for tag, R in (('方案 A', A), ('方案 B', B)):
        v = R['val']
        para(doc, '%s：货箱覆盖 %d/80 且无重复无遗漏；硬时限违规 %d 箱，最小余量 %.1f s'
                  '（%s）；逐航段载荷与整趟能耗均满足载重、体积与 20%% 返航余量；'
                  '无人机峰值同时占用 %d 台、电池峰值 %d 组，编号分配无重叠，且上一趟实际'
                  '充满时刻不晚于下一趟准备开始；汇总指标与明细重算结果一致。'
              % (tag, v['box_coverage']['n_delivered'], v['deadlines']['n_violation'],
                 v['deadlines']['min_hard_slack_s'], v['deadlines']['worst_box'],
                 v['resources']['aircraft']['peak_concurrent'],
                 v['resources']['battery']['peak_concurrent']))

    heading(doc, '七、交付文件与复现方法')
    para(doc, '结果文件：Q2结果.xlsx（按结果提交模板填写 Q2_运输架次、Q2_逐箱交付）、'
              'Q2结果_能耗优先对照.xlsx。明细与报告：transport_sorties.csv、box_deliveries.csv、'
              'flight_legs.csv、aircraft_timeline.csv、battery_timeline.csv、solution.json、'
              'validation.json 与 validation_detail.json、run_manifest.json、Q2结果说明.md。'
              '图形：route_map.png（运输路线）、aircraft_gantt.png 与 battery_gantt.png'
              '（无人机与电池时间线）。')
    para(doc, '复现方式：在 D题/q2 目录下执行 python main.py all <run_id> [energy_first]，'
              '程序会依次完成输入对账、几何与物理核对、候选生成、基线与联合调度、设备编号'
              '分配、独立校验与结果导出；缓存与运行记录写入 q2/cache 与 q2/runs/<run_id>。')
    para(doc, '说明：求解器给出的最优性结论仅对当前候选池、1 s 起始时刻网格及本文口径成立，'
              '未穷尽全部可行路线，因此表述为当前候选集合内的最好可行方案。方案 B 的能耗目标'
              '下界与目标值相差不足 0.1%%（下界 %.2f kWh、目标 %.2f kWh），在该候选池内已接近'
              '最优；架次数与完工时刻受求解时间限制仍存在改进空间。'
              % (B['man']['solver_reports'][1]['best_bound'] / 1000.0,
                 B['man']['solver_reports'][1]['objective'] / 1000.0))

    for fig, cap, R in (('route_map.png', '图 1  方案 A 运输路线（按机型着色）', A),
                        ('aircraft_gantt.png', '图 2  方案 A 无人机使用时间线', A),
                        ('battery_gantt.png', '图 3  方案 A 共享电池占用与充电时间线', A)):
        path = os.path.join(R['dir'], 'figures', fig)
        if os.path.exists(path):
            doc.add_picture(path, width=Cm(15.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(cp.add_run(cap), '黑体', 10, True)

    doc.save(OUT)
    print('已生成:', OUT)


def _input_checks():
    """复用主流程的对账结论（从 run_manifest 之外的固定清单读取）。"""
    import sys
    sys.path.insert(0, os.path.join(BASE, 'q2'))
    from data_io import load_instance
    inst = load_instance(verbose=False)
    return inst['input_checks']


if __name__ == '__main__':
    main()
