# -*- coding: utf-8 -*-
"""P0 辅助：打印全部输入附件的工作表、表头与样例行，供规范化字段对照。"""
import os
import openpyxl
import tifffile
import numpy as np

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
D = os.path.join(BASE, '数据', '无人机应急物资运输基础数据')

files = [
    '调度中心与服务区.xlsx',
    '物资需求与配送时限.xlsx',
    '运输无人机数据.xlsx',
    '中继无人机数据.xlsx',
    '通信链路参数.xlsx',
]
for f in files:
    p = os.path.join(D, f)
    print('=' * 100)
    print('FILE:', f)
    wb = openpyxl.load_workbook(p, data_only=True)
    for ws in wb.worksheets:
        print('-' * 90)
        print('  SHEET:', ws.title, 'dims', ws.dimensions, ws.max_row, 'x', ws.max_column)
        n = 0
        for row in ws.iter_rows(values_only=True):
            vals = ['' if v is None else (f'{v:.6g}' if isinstance(v, float) else str(v)) for v in row]
            while vals and vals[-1] == '':
                vals.pop()
            print('   |', ' | '.join(vals))
            n += 1
            if n >= 60:
                print('   ... (truncated)')
                break

print('=' * 100)
print('RESULT TEMPLATE')
p = os.path.join(BASE, '结果提交模板.xlsx')
wb = openpyxl.load_workbook(p, data_only=True)
for ws in wb.worksheets:
    print('-' * 90)
    print('  SHEET:', ws.title, ws.max_row, 'x', ws.max_column)
    n = 0
    for row in ws.iter_rows(values_only=True):
        vals = ['' if v is None else str(v) for v in row]
        print('   |', ' | '.join(vals))
        n += 1
        if n >= 12:
            print('   ... (truncated)')
            break

print('=' * 100)
print('DEM TAGS')
dem_p = os.path.join(BASE, '数据', '镇龙乡地理空间数据', '镇龙乡及周边地理数据',
                     '数字高程模型数据（DEM）', '镇龙乡及周边30米DEM.tif')
with tifffile.TiffFile(dem_p) as tf:
    pg = tf.pages[0]
    print('shape', pg.shape, 'dtype', pg.dtype)
    for tag in pg.tags.values():
        name = tag.name
        if name in ('ModelPixelScaleTag', 'ModelTiepointTag', 'ModelTransformationTag',
                    'GeoKeyDirectoryTag', 'GeoDoubleParamsTag', 'GeoAsciiParamsTag',
                    'GDAL_NODATA', 'GeoAsciiParamsTag'):
            try:
                v = tag.value
            except Exception as ex:
                v = f'<{ex}>'
            print('  ', name, '=', v)
    a = np.asarray(tf.asarray(), dtype=np.float64)
    print('min/max', np.nanmin(a), np.nanmax(a))
    print('unique-low', np.unique(a)[:5], 'nan count', int(np.isnan(a).sum()))
