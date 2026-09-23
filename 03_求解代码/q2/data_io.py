# -*- coding: utf-8 -*-
"""P0：读取并规范化附件输入，做输入对账。

对外接口：load_instance() -> dict，包含 nodes / boxes / types / aircraft / batteries /
summary；其中质量用整数克、体积用整数立方厘米保存，避免 0.067 m³ 这类临界装载因
浮点误差误判，能耗计算时再换算回 kg。
"""
import os
import json
import hashlib
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
Q2_DIR = os.path.join(BASE_DIR, 'q2')
DATA_DIR = os.path.join(BASE_DIR, '数据')
BASIC_DIR = os.path.join(DATA_DIR, '无人机应急物资运输基础数据')
DEM_DIR = os.path.join(DATA_DIR, '镇龙乡地理空间数据', '镇龙乡及周边地理数据',
                       '数字高程模型数据（DEM）')
DEM_TIF = os.path.join(DEM_DIR, '镇龙乡及周边30米DEM.tif')
TEMPLATE = os.path.join(BASE_DIR, '结果提交模板.xlsx')
RUN_DIR = os.path.join(Q2_DIR, 'runs')
CACHE_DIR = os.path.join(Q2_DIR, 'cache')


def load_config():
    with open(os.path.join(Q2_DIR, 'config.json'), encoding='utf-8') as f:
        return json.load(f)


def _num(v):
    return None if v is None or v == '' else float(v)


def load_nodes():
    p = os.path.join(BASIC_DIR, '调度中心与服务区.xlsx')
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb['数据']
    nodes = {}
    for row in ws.iter_rows(values_only=True):
        if not row or not row[0] or not isinstance(row[0], str):
            continue
        key = row[0].strip()
        if key == 'O01':
            nodes['O01'] = dict(node_id='O01', kind='depot', name=row[1],
                                lon=float(row[2]), lat=float(row[3]), ground_z=float(row[4]),
                                population=None)
        elif len(key) == 4 and key[0] == 'S' and key[1:].isdigit():
            nodes[key] = dict(node_id=key, kind='zone', name=row[1],
                              lon=float(row[2]), lat=float(row[3]), ground_z=float(row[4]),
                              population=_num(row[5]) if len(row) > 5 else None)
    for nid, n in nodes.items():
        n['work_z'] = n['ground_z'] if nid == 'O01' else n['ground_z'] + 30.0
    return nodes


def load_types():
    p = os.path.join(BASIC_DIR, '运输无人机数据.xlsx')
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb['数据']
    keys = ['type_id', 'type_name', 'mass_kg', 'cap_mass_kg', 'cap_volume_m3', 'v_cruise',
            'range_empty_m', 'range_full_m', 'energy_kwh', 'reserve_pct', 't_prep_s',
            't_load_s', 't_hand_base_s', 't_hand_box_s', 'v_up', 'v_down',
            'eta_up', 'eta_down']
    types = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0] in ('A', 'B', 'C') and isinstance(row[1], str) and len(row[1]) > 4:
            types[row[0]] = dict(zip(keys, row))
    for g, t in types.items():
        t['energy_usable_kwh'] = t['energy_kwh'] * (1 - t['reserve_pct'] / 100.0)
    return types


def load_aircraft():
    p = os.path.join(BASIC_DIR, '运输无人机数据.xlsx')
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb['数据']
    aircraft = {}
    for row in ws.iter_rows(values_only=True):
        if row and isinstance(row[0], str) and row[0].startswith('U') and len(row[0]) == 3:
            aircraft[row[0]] = dict(aircraft_id=row[0], type_id=row[1], start=row[2])
    return aircraft


def load_battery_stock():
    """附件口径：共享电池组总数已包含初始装机电池，不能额外再加。"""
    p = os.path.join(BASIC_DIR, '运输无人机数据.xlsx')
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb['数据']
    stock = {}
    for row in ws.iter_rows(values_only=True):
        if (row and row[0] in ('A', 'B', 'C') and isinstance(row[1], (int, float))
                and isinstance(row[2], (int, float))):
            stock[row[0]] = dict(count=int(row[1]), full_charge_s=float(row[2]))
    return stock


def load_boxes(nodes):
    p = os.path.join(BASIC_DIR, '物资需求与配送时限.xlsx')
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb['逐箱货箱清单']
    boxes = []
    for row in ws.iter_rows(values_only=True):
        if not row or not isinstance(row[0], str) or '-' not in row[0]:
            continue
        box_id, zone, cat, mass, vol, first, first_dl, exp, prio = row[:9]
        first = (str(first).strip() == '是')
        first_dl = _num(first_dl)
        exp = _num(exp)
        # 硬时限：医疗物资取期望送达时间，首批保障货箱取首批截止时间，同时适用取较早者
        cand = []
        if cat == '医疗物资' and exp is not None:
            cand.append(exp)
        if first and first_dl is not None:
            cand.append(first_dl)
        hard = min(cand) if cand else None
        boxes.append(dict(
            box_id=box_id.strip(), zone_id=zone.strip(), category=cat.strip(),
            mass_kg=float(mass), volume_m3=float(vol),
            mass_g=int(round(float(mass) * 1000)), vol_cc=int(round(float(vol) * 1e6)),
            first_batch=first, first_deadline_s=first_dl, expected_s=exp,
            priority=float(prio), hard_deadline_s=hard))
    return boxes


def load_demand_sheet():
    p = os.path.join(BASIC_DIR, '物资需求与配送时限.xlsx')
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb['数据']
    out = []
    for row in ws.iter_rows(values_only=True):
        if row and isinstance(row[0], str) and len(row[0]) == 4 and row[0][0] == 'S' \
                and row[0][1:].isdigit():
            out.append(dict(zone_id=row[0], category=row[1], boxes=int(row[2]),
                            first_boxes=int(row[3]), mass_kg=float(row[4]),
                            volume_m3=float(row[5]), priority=float(row[6]),
                            first_deadline_s=_num(row[7]), expected_s=_num(row[8])))
    return out


def load_instance(verbose=True):
    nodes = load_nodes()
    types = load_types()
    aircraft = load_aircraft()
    stock = load_battery_stock()
    boxes = load_boxes(nodes)

    batteries = {}
    for g, s in stock.items():
        for k in range(1, s['count'] + 1):
            bid = 'B%s%02d' % (g, k)
            batteries[bid] = dict(battery_id=bid, type_id=g, soc0=1.0,
                                  full_charge_s=s['full_charge_s'])
        if g in types:
            types[g]['full_charge_s'] = s['full_charge_s']
            types[g]['n_batteries'] = s['count']
    inst = dict(nodes=nodes, types=types, aircraft=aircraft,
                battery_stock=stock, batteries=batteries, boxes=boxes)

    # ---- 输入对账 ----
    checks = []
    zones = [n for n in nodes if n != 'O01']
    checks.append(('节点数 = 16（1 调度中心 + 15 服务区）', len(nodes) == 16 and len(zones) == 15,
                   '%d 个节点' % len(nodes)))
    checks.append(('货箱数 = 80 且编号唯一', len(boxes) == 80 and len({b['box_id'] for b in boxes}) == 80,
                   '%d 箱 / %d 个唯一编号' % (len(boxes), len({b['box_id'] for b in boxes}))))
    tm = sum(b['mass_g'] for b in boxes)
    tv = sum(b['vol_cc'] for b in boxes)
    checks.append(('总质量 = 758 kg', tm == 758000, '%.3f kg' % (tm / 1000)))
    checks.append(('总体积 = 2.011 m³', abs(tv - 2011000) <= 2000, '%.4f m³' % (tv / 1e6)))
    fb = [b for b in boxes if b['first_batch']]
    hd = [b for b in boxes if b['hard_deadline_s'] is not None]
    checks.append(('首批保障货箱 = 30 箱', len(fb) == 30, '%d 箱' % len(fb)))
    checks.append(('硬时限货箱 = 31 箱', len(hd) == 31, '%d 箱' % len(hd)))
    for dl, expect in ((3600.0, 17), (7200.0, 7), (10800.0, 7)):
        cnt = len([b for b in hd if b['hard_deadline_s'] == dl])
        checks.append(('硬截止 %.0f s 对应 %d 箱' % (dl, expect), cnt == expect, '%d 箱' % cnt))
    b02 = next(b for b in boxes if b['box_id'] == 'S001-MED-02')
    checks.append(('S001-MED-02 硬截止 = 3600 s（非首批但须按时）', b02['hard_deadline_s'] == 3600.0,
                   str(b02['hard_deadline_s'])))
    for bid in ('S012-MED-01', 'S014-MED-01'):
        b = next(x for x in boxes if x['box_id'] == bid)
        checks.append(('%s 取较早时限 3600 s' % bid, b['hard_deadline_s'] == 3600.0,
                       str(b['hard_deadline_s'])))
    # 汇总需求与逐箱清单逐区逐类型对账
    demo = load_demand_sheet()
    agg = {}
    for b in boxes:
        k = (b['zone_id'], b['category'])
        a = agg.setdefault(k, [0, 0, 0, 0, None, set()])
        a[0] += 1
        a[1] += 1 if b['first_batch'] else 0
        a[2] += b['mass_g']
        a[3] += b['vol_cc']
        a[4] = b['priority']
        a[5].add(b['expected_s'])
    mismatch = []
    for d in demo:
        k = (d['zone_id'], d['category'])
        a = agg.get(k)
        if a is None:
            mismatch.append((k, 'missing'))
            continue
        # 汇总表给的是“总需求箱数/首批箱数”与“单箱质量/单箱体积/应急优先系数”
        ok = (a[0] == d['boxes'] and a[1] == d['first_boxes']
              and abs(a[2] / a[0] - d['mass_kg'] * 1000) <= 1
              and abs(a[3] / a[0] - d['volume_m3'] * 1e6) <= 1
              and abs(a[4] - d['priority']) < 1e-9)
        if not ok:
            mismatch.append((k, a[:5], d))
    checks.append(('汇总需求与逐箱清单逐区逐类型一致', not mismatch, str(mismatch if mismatch else 'OK')))

    inst['summary'] = dict(
        n_nodes=len(nodes), n_boxes=len(boxes), total_mass_kg=tm / 1000.0,
        total_volume_m3=tv / 1e6, n_first_batch=len(fb), n_hard=len(hd),
        hard_by_deadline={str(int(dl)): len([b for b in hd if b['hard_deadline_s'] == dl])
                          for dl in (3600.0, 7200.0, 10800.0)},
        zones_with_hard_first_hour=sorted({b['zone_id'] for b in hd
                                           if b['hard_deadline_s'] == 3600.0}),
        total_energy_if_direct=None)
    inst['input_checks'] = checks
    if verbose:
        print('[P0] 输入对账')
        for name, ok, info in checks:
            print('   %-42s %-4s %s' % (name, 'PASS' if ok else 'FAIL', info))
        allok = all(c[1] for c in checks)
        print('   => %s' % ('全部通过' if allok else '存在失败项，请先修正输入口径'))
    if not all(c[1] for c in checks):
        raise SystemExit('输入对账失败，终止求解。')
    return inst


def input_hash():
    h = hashlib.sha256()
    for p in [os.path.join(BASIC_DIR, f) for f in sorted(os.listdir(BASIC_DIR))]:
        if p.endswith('.xlsx') and not os.path.basename(p).startswith('~$'):
            h.update(os.path.basename(p).encode('utf-8'))
            with open(p, 'rb') as f:
                h.update(f.read())
    with open(DEM_TIF, 'rb') as f:
        h.update(f.read(65536))
    return h.hexdigest()[:16]


if __name__ == '__main__':
    inst = load_instance()
    print('输入哈希:', input_hash())
